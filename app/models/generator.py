import os
import uuid
from docx import Document

def generate_loan_agreement(data):
    template_path = os.path.join('templates', 'loan_agreement_base.docx')
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found at {template_path}")
    
    doc = Document(template_path)
    
    # Replace placeholders in paragraphs
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            placeholder = f"{{{{ {key} }}}}"
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, str(value))
                
    # Check tables if any placeholders are there
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for key, value in data.items():
                        placeholder = f"{{{{ {key} }}}}"
                        if placeholder in paragraph.text:
                            paragraph.text = paragraph.text.replace(placeholder, str(value))

    generated_id = str(uuid.uuid4())
    filename = f"loan_{generated_id}.docx"
    output_dir = os.environ.get("GENERATED_FOLDER", "static/generated")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, filename)
    
    doc.save(output_path)
    
    # Ensure forward slashes for URL download path
    normalized_output_dir = output_dir.replace('\\', '/')
    download_url = f"/{normalized_output_dir}/{filename}"
    
    # Generate HTML string preview
    preview_html = "<html><body><h2>LOAN AGREEMENT</h2>"
    preview_html += f"<p>This Loan Agreement ('Agreement') is entered into on {data.get('start_date', '')}, by and between:</p>"
    preview_html += f"<p>Borrower: {data.get('borrower_name', '')}</p>"
    preview_html += f"<p>Lender: {data.get('lender_name', '')}</p>"
    preview_html += "<h3>1. Loan Amount and Interest</h3>"
    preview_html += f"<p>The Lender agrees to loan the Borrower the principal sum of {data.get('loan_amount', '')} ('Loan'), together with interest on the outstanding principal amount at the rate of {data.get('interest_rate', '')}% per annum.</p>"
    preview_html += "<h3>2. Repayment</h3>"
    preview_html += f"<p>The Loan shall be repaid in full over a tenure of {data.get('tenure_months', '')} months.</p>"
    preview_html += "<h3>3. Collateral</h3>"
    preview_html += f"<p>The Borrower agrees to pledge the following collateral as security for the Loan: {data.get('collateral', '')}</p>"
    preview_html += "<h3>4. Penalty for Default</h3>"
    preview_html += f"<p>In the event of default, the following penalty clause shall apply: {data.get('penalty_clause', '')}</p>"
    preview_html += "<h3>5. Governing Law</h3>"
    preview_html += f"<p>This Agreement shall be governed by and construed in accordance with the laws of: {data.get('governing_law', '')}.</p>"
    preview_html += "</body></html>"

    return {
        "document_id": generated_id,
        "download_url": download_url,
        "preview_html": preview_html
    }
